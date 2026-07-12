"""Extraction de texte brut à partir d'un CV PDF ou DOCX (MVP — pas d'OCR)."""

from pathlib import Path

FORMATS_SUPPORTES = {".pdf", ".docx"}


def extraire_texte(chemin: Path) -> str:
    suffixe = chemin.suffix.lower()
    if suffixe == ".pdf":
        texte = _extraire_pdf(chemin)
    elif suffixe == ".docx":
        texte = _extraire_docx(chemin)
    else:
        raise ValueError(f"Format non supporté : {suffixe!r} (attendu : {sorted(FORMATS_SUPPORTES)})")

    if not texte.strip():
        raise ValueError(
            f"Aucun texte exploitable extrait de {chemin} (PDF scanné sans OCR ? document vide ?)"
        )
    return texte


def _extraire_pdf(chemin: Path) -> str:
    import pdfplumber

    with pdfplumber.open(chemin) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(pages).strip()


def _extraire_docx(chemin: Path) -> str:
    """Beaucoup de CV utilisent des tableaux Word pour la mise en page (bloc nom/contact,
    grille de compétences) — `document.paragraphs` seul les ignore silencieusement."""
    import docx

    document = docx.Document(str(chemin))
    morceaux = [p.text for p in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cellules_vues: set[int] = set()
            textes_ligne = []
            for cell in row.cells:
                identifiant_cellule = id(cell._tc)  # cellules fusionnées répétées sinon
                if identifiant_cellule in cellules_vues:
                    continue
                cellules_vues.add(identifiant_cellule)
                if cell.text.strip():
                    textes_ligne.append(cell.text.strip())
            if textes_ligne:
                morceaux.append(" | ".join(textes_ligne))

    return "\n".join(m for m in morceaux if m.strip()).strip()
