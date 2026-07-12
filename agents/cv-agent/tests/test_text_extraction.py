from pathlib import Path

import docx
import pytest
from fpdf import FPDF

from cv_agent.text_extraction import extraire_texte

LIGNES_CV = ["Jean Dupont", "Développeur Backend Python", "Expérience : Acme Corp, 2020-2023"]


def _creer_pdf(chemin: Path, lignes: list[str]) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for ligne in lignes:
        pdf.cell(0, 10, text=ligne, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(chemin))


def _creer_docx(chemin: Path, lignes: list[str]) -> None:
    document = docx.Document()
    for ligne in lignes:
        document.add_paragraph(ligne)
    document.save(str(chemin))


def test_extraire_texte_pdf(tmp_path: Path) -> None:
    chemin = tmp_path / "cv.pdf"
    _creer_pdf(chemin, LIGNES_CV)

    texte = extraire_texte(chemin)

    assert "Jean Dupont" in texte
    assert "Développeur Backend Python" in texte


def test_extraire_texte_docx(tmp_path: Path) -> None:
    chemin = tmp_path / "cv.docx"
    _creer_docx(chemin, LIGNES_CV)

    texte = extraire_texte(chemin)

    assert "Jean Dupont" in texte
    assert "Acme Corp" in texte


def test_extraire_texte_format_non_supporte(tmp_path: Path) -> None:
    chemin = tmp_path / "cv.txt"
    chemin.write_text("contenu", encoding="utf-8")

    with pytest.raises(ValueError, match="non supporté"):
        extraire_texte(chemin)


def test_extraire_texte_pdf_vide_leve_erreur(tmp_path: Path) -> None:
    chemin = tmp_path / "vide.pdf"
    _creer_pdf(chemin, [])

    with pytest.raises(ValueError, match="Aucun texte"):
        extraire_texte(chemin)
